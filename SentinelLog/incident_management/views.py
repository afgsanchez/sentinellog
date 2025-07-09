from django.shortcuts import render, redirect, get_object_or_404
from django.db.models import Q
from .models import *
from .forms import IncidentForm, IncidentNoteForm
from docx import Document
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone
from fpdf import FPDF
import os
from io import BytesIO


def extract_docx_text(docx_file):
    try:
        document = Document(docx_file)
        text = "\n".join([para.text for para in document.paragraphs])
        return text
    except Exception:
        return ""

def index(request):
    incidents = Incident.objects.all().order_by('-date')
    search = request.GET.get('search')
    if search:
        incidents = incidents.filter(
            Q(description__icontains=search) |
            Q(location__icontains=search) |
            Q(insurance_case_number__icontains=search) |
            Q(docx_text__icontains=search)
        )
    return render(request, 'incident_management/index.html', {
        'incidents': incidents,
        'search': search,
    })

def add_incident(request):
    if request.method == 'POST':
        form = IncidentForm(request.POST, request.FILES)
        if form.is_valid():
            incident = form.save(commit=False)
            docx_file = request.FILES.get('incident_report_docx')
            if docx_file:
                incident.docx_text = extract_docx_text(docx_file)
            incident.save()
            return redirect('incident_management_index')
    else:
        form = IncidentForm()
    return render(request, 'incident_management/add_incident.html', {'form': form})

def edit_incident(request, pk):
    incident = get_object_or_404(Incident, pk=pk)
    if request.method == 'POST':
        form = IncidentForm(request.POST, request.FILES, instance=incident)
        if form.is_valid():
            incident = form.save(commit=False)
            docx_file = request.FILES.get('incident_report_docx')
            if docx_file:
                incident.docx_text = extract_docx_text(docx_file)
            incident.save()
            return redirect('incident_management_index')
    else:
        form = IncidentForm(instance=incident)
    return render(request, 'incident_management/edit_incident.html', {'form': form, 'incident': incident})

from .models import Incident, IncidentPhoto, IncidentNote, IncidentAttachment

from .forms import IncidentPhotoForm, IncidentAttachmentForm

@login_required
def incident_detail(request, pk):
    incident = get_object_or_404(Incident, pk=pk)
    photos = incident.photos.all()
    notes = incident.notes.all().order_by('-created_at')
    attachments = incident.attachments.all()

    note_form = IncidentNoteForm()
    photo_form = IncidentPhotoForm()
    attach_form = IncidentAttachmentForm()

    if request.method == 'POST':
        if 'add_note' in request.POST:
            note_form = IncidentNoteForm(request.POST)
            if note_form.is_valid():
                note = note_form.save(commit=False)
                note.incident = incident
                note.author = request.user
                note.save()
                return redirect('incident_detail', pk=pk)
        elif 'add_photo' in request.POST:
            photo_form = IncidentPhotoForm(request.POST, request.FILES)
            if photo_form.is_valid():
                photo = photo_form.save(commit=False)
                photo.incident = incident
                photo.uploaded_by = request.user
                photo.save()
                return redirect('incident_detail', pk=pk)
        elif 'add_attachment' in request.POST:
            attach_form = IncidentAttachmentForm(request.POST, request.FILES)
            if attach_form.is_valid():
                attach = attach_form.save(commit=False)
                attach.incident = incident
                attach.uploaded_by = request.user
                attach.save()
                return redirect('incident_detail', pk=pk)

    return render(request, 'incident_management/incident_detail.html', {
        'incident': incident,
        'photos': photos,
        'notes': notes,
        'attachments': attachments,
        'note_form': note_form,
        'photo_form': photo_form,
        'attach_form': attach_form,
    })

def generate_incident_pdf(request, pk):
    incident = get_object_or_404(Incident, pk=pk)
    photos = incident.photos.all()
    notes = incident.notes.all().order_by('-created_at')
    attachments = incident.attachments.all()

    class CyberReportPDF(FPDF):
        def header(self):
            # Si quieres probar sin logo, comenta la siguiente línea
            logo_path = os.path.join(os.path.dirname(__file__), 'static', 'SentinelLog_logo.png')
            self.image(logo_path, 10, 8, 33)
            self.set_font('DejaVu', 'B', 16)
            self.set_text_color(30, 30, 30)
            self.cell(0, 12, 'MVC Son Antem | Safety & Security', ln=True, align='C')
            self.set_font('DejaVu', '', 12)
            self.set_text_color(100, 100, 100)
            self.cell(0, 8, 'Reporte confidencial', ln=True, align='C')
            self.ln(5)
            self.set_draw_color(0, 102, 204)
            self.set_line_width(1)
            self.line(10, self.get_y(), 200, self.get_y())
            self.ln(8)

        def footer(self):
            self.set_y(-18)
            self.set_font('DejaVu', 'I', 8)
            self.set_text_color(120, 120, 120)
            self.cell(0, 8, f'Página {self.page_no()}', align='L')
            self.cell(0, 8, f'Generado el: {timezone.localtime(timezone.now()).strftime("%d/%m/%Y %H:%M")}', align='R')
        
        def section_title(self, title):
            self.set_font('DejaVu', 'B', 14)
            self.set_text_color(0, 102, 204)
            self.cell(0, 10, title, ln=True)
            self.set_text_color(0, 0, 0)
            self.set_font('DejaVu', '', 12)
            self.ln(2)

        def add_divider(self):
            self.set_draw_color(200, 200, 200)
            self.set_line_width(0.5)
            self.line(10, self.get_y(), 200, self.get_y())
            self.ln(5)

    pdf = CyberReportPDF()

    font_dir = os.path.join(os.path.dirname(__file__), 'static', 'fonts')
    pdf.add_font('DejaVu', '', os.path.join(font_dir, 'DejaVuSans.ttf'), uni=True)
    pdf.add_font('DejaVu', 'B', os.path.join(font_dir, 'DejaVuSans-Bold.ttf'), uni=True)
    pdf.add_font('DejaVu', 'I', os.path.join(font_dir, 'DejaVuSans-Oblique.ttf'), uni=True)
    pdf.add_font('DejaVu', 'BI', os.path.join(font_dir, 'DejaVuSans-BoldOblique.ttf'), uni=True)

    pdf.add_page()

    # Portada
    pdf.set_font("DejaVu", 'B', 20)
    pdf.set_text_color(0, 51, 102)
    pdf.cell(0, 60, '', ln=True)  # Espacio
    pdf.cell(0, 10, f"Incident Report:", ln=True, align='C')
    pdf.set_font("DejaVu", '', 14)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(0, 10, f"Persona afectada: {incident.affected_person or '-'}", ln=True, align='C')
    pdf.cell(0, 10, f"Ubicación: {incident.location}", ln=True, align='C')
    pdf.cell(0, 10, f"Fecha y hora: {incident.date.strftime('%d/%m/%Y %H:%M')}", ln=True, align='C')
    pdf.ln(20)
    pdf.set_font("DejaVu", 'I', 12)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(0, 10, "Informe generado automáticamente por SentinelLog", ln=True, align='C')
    pdf.add_page()

    # Descripción
    pdf.section_title("1. Descripción del Incidente")
    pdf.set_font("DejaVu", '', 12)
    pdf.multi_cell(0, 8, incident.description or "")
    pdf.add_divider()

    # Datos principales
    pdf.section_title("2. Datos principales")
    pdf.set_font("DejaVu", '', 12)
    pdf.cell(0, 8, f"Reportado por: {incident.reported_by}", ln=True)
    pdf.cell(0, 8, f"Caso aseguradora: {incident.insurance_case_number or '-'}", ln=True)
    pdf.cell(0, 8, f"Investigación relacionada: {incident.related_investigation.title if incident.related_investigation else '-'}", ln=True)
    pdf.add_divider()

    # Notas
    pdf.section_title("3. Notas")
    if notes:
        for note in notes:
            pdf.set_font("DejaVu", 'B', 11)
            pdf.set_text_color(0, 102, 204)
            pdf.cell(0, 7, f"{note.author} ({note.created_at.strftime('%d/%m/%Y %H:%M')})", ln=True)
            pdf.set_font("DejaVu", '', 11)
            pdf.set_text_color(0, 0, 0)
            pdf.multi_cell(0, 7, note.note or "")
            pdf.ln(2)
    else:
        pdf.set_font("DejaVu", 'I', 11)
        pdf.cell(0, 8, "No hay notas registradas.", ln=True)
    pdf.add_divider()

    # Archivos adjuntos
    pdf.section_title("4. Archivos adjuntos")
    if attachments:
        for attach in attachments:
            pdf.set_font("DejaVu", 'B', 11)
            pdf.set_text_color(0, 123, 255)
            pdf.cell(0, 7, f"{attach.file.name.split('/')[-1]}", ln=True)
            pdf.set_font("DejaVu", '', 11)
            pdf.set_text_color(0, 0, 0)
            pdf.cell(0, 7, f"Descripción: {attach.description or '-'}", ln=True)
            pdf.cell(0, 7, f"Subido por: {attach.uploaded_by} el {attach.uploaded_at.strftime('%d/%m/%Y %H:%M')}", ln=True)
            pdf.ln(2)
    else:
        pdf.set_font("DejaVu", 'I', 11)
        pdf.cell(0, 8, "No hay archivos adjuntos.", ln=True)
    pdf.add_divider()

    # Fotos (solo nombres, no imágenes)
    pdf.section_title("5. Fotos adjuntas")
    if photos:
        for photo in photos:
            pdf.cell(0, 7, f"{photo.image.name.split('/')[-1]} (Subida por: {photo.uploaded_by} el {photo.uploaded_at.strftime('%d/%m/%Y %H:%M')})", ln=True)
    else:
        pdf.set_font("DejaVu", 'I', 11)
        pdf.cell(0, 8, "No hay fotos adjuntas.", ln=True)

    pdf_bytes = pdf.output(dest='S')
    if isinstance(pdf_bytes, bytearray):
        pdf_bytes = bytes(pdf_bytes)
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename=incidente_{incident.pk}_reporte.pdf'
    return response
    

# def test_pdf(request):
#     from fpdf import FPDF
#     pdf = FPDF()
#     pdf.add_page()
#     pdf.set_font("Arial", size=12)
#     pdf.cell(200, 10, txt="¡Hola, PDF!", ln=True, align='C')
#     pdf_bytes = pdf.output(dest='S')
#     if isinstance(pdf_bytes, str):
#         pdf_bytes = pdf_bytes.encode('latin1')
#     elif isinstance(pdf_bytes, bytearray):
#         pdf_bytes = bytes(pdf_bytes)
#     return HttpResponse(pdf_bytes, content_type='application/pdf')